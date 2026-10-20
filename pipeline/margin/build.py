# -*- coding: utf-8 -*-
"""月更建置：FINRA xlsx + yahoo_raw.json → data.json → 網站 margin.html（模板 template.html）
本機另寫 dashboard.html（Artifact 用）並更新疊圖數據 xlsx（路徑存在才寫）。
用法：python build.py          手動月更（不顯示提示條）
      python build.py --auto   GitHub Actions 自動更新（頁首加「文字說明待更新」提示條）
      只改版面不更新數據時：MARGIN_UPDATED=2026/10/06 python build.py（頁尾日期不變）"""
import json, math, sys, io, os, datetime as dt
import pandas as pd
os.chdir(os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding="utf-8")
AUTO = "--auto" in sys.argv
UPDATED = (dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=8)).strftime("%Y/%m/%d")   # 台灣日期
UPDATED = os.environ.get("MARGIN_UPDATED", UPDATED)   # 只改版面重產時沿用原「最後更新」日期

fin = pd.read_excel("margin-statistics.xlsx")
fin.columns = ["ym","mar","fc_cash","fc_margin"]
fin = fin[["ym","mar"]].dropna()
fin["ym"] = fin["ym"].astype(str).str[:7]
mar_map = dict(zip(fin.ym, fin.mar.astype(int)))
ya = json.load(open("yahoo_raw.json"))

last_idx_month = max(ya["spx"])            # 指數最新月（可能為月中值）
last_fin_month = max(mar_map)
months=[]; y,m=1997,1
while f"{y}-{m:02d}" <= last_idx_month:
    months.append(f"{y}-{m:02d}"); m+=1
    if m==13: y+=1; m=1
N=len(months); idx={mo:i for i,mo in enumerate(months)}

def series(k, nd=1):
    return [round(ya[k][mo],nd) if mo in ya[k] and ya[k][mo] is not None else None for mo in months]
mar=[mar_map.get(mo) for mo in months]
yoy=[None]*N
for i in range(12,N):
    if mar[i] and mar[i-12]: yoy[i]=round((mar[i]/mar[i-12]-1)*100,1)
spx=series("spx"); ndx=series("ndx"); sox=series("sox")
# W5000 對數插值補缺
w=[ya["w5000"].get(mo) for mo in months]
known=[i for i,v in enumerate(w) if v]
for i in range(N):
    if w[i]: continue
    lo=max([k for k in known if k<i],default=None); hi=min([k for k in known if k>i],default=None)
    if lo is not None and hi is not None:
        w[i]=math.exp(math.log(w[lo])+(math.log(w[hi])-math.log(w[lo]))*(i-lo)/(hi-lo))
    elif lo is not None: w[i]=w[lo]
ratio=[round(mar[i]/(w[i]*1050)*100,3) if mar[i] and w[i] else None for i in range(N)]
data={"months":months,"mar":mar,"yoy":yoy,"spx":spx,"ndx":ndx,"sox":sox,"ratio":ratio}
# IBKR nowcast（手動維護 nowcast.json：month/yoy/lo/hi/label/note；紅色虛線自最後一個實際月延伸到 month）
import os
if os.path.exists("nowcast.json"):
    ncs=json.load(open("nowcast.json",encoding="utf-8"))
    if isinstance(ncs,dict): ncs=[ncs]          # 舊格式單筆相容；新格式為清單，歷次暫估全保留供對帳
    data["nowcast"]=[]
    for nc in ncs:
        if nc.get("month") not in idx:
            print("nowcast",nc.get("month"),"超出 months 範圍，略過"); continue
        data["nowcast"].append(nc)
        act=yoy[idx[nc["month"]]]
        print("nowcast",nc["month"],nc["yoy"],"| FINRA 實際",act,"| 誤差",None if act is None else round(nc["yoy"]-act,1),"pp")
# 頁首燈號③：最近一次 FOMC 決議（手動維護 fed.json：date/action=升息|維持|降息|量化寬鬆/text）
data["fed"]=json.load(open("fed.json",encoding="utf-8"))
json.dump(data,open("data.json","w",encoding="utf-8"),ensure_ascii=False)
print("months",months[0],"→",months[-1],"| FINRA 最新",last_fin_month, mar_map[last_fin_month], "YoY",yoy[idx[last_fin_month]])
# YoY 恆等式（下月）
nm_i=idx[last_fin_month]+1
if nm_i<N:
    base=mar[nm_i-12]; a=mar[nm_i-1]/base-1; b=mar[nm_i-1]/base
    print(f"下月 YoY 恆等式: YoY = {a*100:.1f}% + {b:.3f} × MoM ；站回40%需 MoM {(1.40/b-1)*100:+.1f}% ；MoM +0.6% → {(a+b*0.006)*100:.1f}%")

# ---- 回測表 ----
# 排序：峰值 <50% 淺觸 → 命中 → 當前事件（置底、黃底 act）；第 5 欄 = tr class
EV=[("1998/4","1998-04",42.5,"淺觸",False,""),
    ("2010/4","2010-04",42.1,"淺觸",False,""),
    ("2025/10","2025-10",45.2,"淺觸",False,""),
    ("1999/10–2000/9","2000-03",80.5,"命中",True,"hit"),
    ("2007/5–10","2007-07",62.6,"命中",True,"hit"),
    ("2021/1–10","2021-03",71.6,"命中",True,"hit"),
    ("<b>2026/4–6</b>","2026-05",53.7,"<b>當前事件</b>",True,"act")]
def cell(v,gl=False,star=False):
    cls=("gl " if gl else "")+("" if v is None else ("pos" if v>=0 else "neg"))
    txt="—" if v is None else f"{v:+.1f}%"+("＊" if star else "")
    return f'<td class="{cls.strip()}">{txt}</td>' if cls.strip() else f"<td>{txt}</td>"
rows=[]; xl_rows=[]
for lab,T,pk,verdict,hit,trc in EV:
    t=idx[T]; ongoing = t+24 > N-1
    tds=f"<td>{lab}</td><td>{'<b>%.1f%%</b>'%pk if hit else '%.1f%%'%pk}</td><td>{verdict}</td>"
    rec={}
    for k,arr in (("spx",spx),("ndx",ndx),("sox",sox)):
        base=arr[t]; vals=[]
        for n in (3,6,12):
            v=round((arr[t+n]/base-1)*100,1) if t+n<N else None; vals.append(v)
        end=min(t+24,N-1); pk_=arr[t]; dd=0.0
        for v in arr[t:end+1]:
            pk_=max(pk_,v); dd=min(dd,(v/pk_-1)*100)
        dd=round(dd,1)
        tds+=cell(vals[0],gl=True)+cell(vals[1])+cell(vals[2])+cell(dd,star=ongoing)
        rec[k]=(vals,dd)
    rows.append(f'      <tr{f" class=\"{trc}\"" if trc else ""}>{tds}</tr>')
    xl_rows.append((lab.replace("<b>","").replace("</b>",""),T,pk,rec))
    print(lab.replace("<b>","").replace("</b>",""), {k:(v[0],v[1]) for k,v in rec.items()})
tpl=open("template.html",encoding="utf-8").read()
html=(tpl.replace("__DATA__",json.dumps(data,ensure_ascii=False))
         .replace("__BACKTEST_ROWS__","\n".join(rows))
         .replace("__LAST_MONTH__",last_idx_month.replace("-","/"))
         .replace("__UPDATED__",UPDATED))
# 頁尾：指數最新月若為抓取當月 → 標註月中值日期
fetched = ya.get("_fetched")
idx_note = f" 為 {int(fetched[5:7])}/{int(fetched[8:10])} 月中值" if fetched and fetched[:7]==last_idx_month else ""
last_fin = f"{last_fin_month[:4]}/{int(last_fin_month[5:])}"   # 2026/8（不補零）
auto_note = (f'\n      <p class="auto-note">⚠ 本頁已於 {UPDATED} 自動更新 FINRA 數據至 {last_fin}：圖表、槓桿熱度表與前兩個燈號已更新；'
             f'文字說明（各表註解、IBKR 表、Fed 救援空間表）尚待人工更新。</p>') if AUTO else ""
html=(html.replace("__AUTO_NOTE__",auto_note).replace("__LAST_FIN__",last_fin).replace("__IDX_NOTE__",idx_note))
for tag in ("__DATA__","__BACKTEST_ROWS__","__LAST_MONTH__","__UPDATED__","__AUTO_NOTE__","__LAST_FIN__","__IDX_NOTE__"):
    assert tag not in html, tag
open("dashboard.html","w",encoding="utf-8").write(html)
print("dashboard.html", len(html), "bytes")

# ---- 宏觀監控儀表板（GitHub Pages）子頁面：補完整文件骨架＋共用目錄欄 nav.js ----
SITE = "../../margin.html"
site_html = ('<!doctype html>\n<html lang="zh-Hant">\n<head>\n<meta charset="utf-8">\n'
             '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
             + html.replace("<title>美國融資槓桿監測</title>", "<title>融資槓桿｜宏觀監控儀表板</title>", 1)
                   .replace('<div class="wrap">', '<style>body{padding-top:0}.wrap{margin-top:20px}</style>\n'
                            '</head>\n<body>\n<script src="nav.js?v=20261020"></script>\n<div class="wrap">', 1)
             + "\n</body>\n</html>\n")
assert site_html.count("nav.js") == 1
open(SITE, "w", encoding="utf-8", newline="\n").write(site_html)
print(SITE, len(site_html), "bytes")

# ---- xlsx（本機才有，CI 略過）----
xp="../../../儀表板開發/融資槓桿疊圖數據.xlsx"
if not os.path.exists(xp):
    print("xlsx 不在，略過"); sys.exit(0)
import openpyxl
wb=openpyxl.load_workbook(xp); ws=wb.worksheets[0]
ws.delete_rows(2, ws.max_row)
for i,mo in enumerate(months):
    ws.append([mo, mar[i], yoy[i], spx[i], ndx[i], sox[i]])
ws2=wb.worksheets[1]; ws2.delete_rows(2, ws2.max_row)
starts={"1998/4":("1998-04","1998-04"),"1999/10–2000/9":("1999-10","2000-09"),"2007/5–10":("2007-05","2007-10"),
        "2010/4":("2010-04","2010-04"),"2021/1–10":("2021-01","2021-10"),"2025/10":("2025-10","2025-10"),"2026/4–6":("2026-04","2026-06")}
for lab,T,pk,rec in xl_rows:
    s,e=starts[lab]; sp=rec["spx"]; so=rec["sox"]
    ws2.append([s,e,T,pk,sp[0][0],sp[0][1],sp[0][2],sp[1],so[0][2],so[1]])
wb.save(xp); print("xlsx updated")
